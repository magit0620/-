from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from quote_builder.__main__ import main
from quote_builder.pricing import PricingRule, price_items
from quote_builder.reader import LineItem, extract_items, parse_number, read_competitor_quote
from quote_builder.writer import QuoteInfo, write_quote

SAMPLES = Path(__file__).resolve().parent.parent / "samples"


@pytest.mark.parametrize("raw, expected", [
    ("¥1,200", 1200), ("１５０，０００円", 150000), ("△500", -500), (3, 3), ("abc", None), ("", None),
])
def test_parse_number(raw, expected):
    assert parse_number(raw) == expected


def test_read_xlsx_sample():
    q = read_competitor_quote(SAMPLES / "competitor_quote.xlsx")
    assert q.vendor == "株式会社ライバル商会"
    assert [i.name for i in q.items][:2] == ["ノートPC", "ディスプレイ 24インチ"]
    assert len(q.items) == 5  # 小計・消費税・合計は除外
    assert q.items[1].unit_price == 23800
    assert q.items[3].unit_price == 150000


def test_read_csv_derives_unit_price_from_amount():
    q = read_competitor_quote(SAMPLES / "competitor_quote.csv")
    lic = next(i for i in q.items if i.name == "ソフトウェアライセンス")
    assert lic.quantity == 25 and lic.unit_price == 15000
    assert len(q.items) == 3


def test_extract_without_header_raises():
    with pytest.raises(ValueError):
        extract_items([["foo", "bar"], [1, 2]])


@pytest.mark.parametrize("mode, expected", [("down", 1130), ("up", 1140), ("nearest", 1140)])
def test_pricing_rounding(mode, expected):
    assert PricingRule(rate=0.95, round_unit=10, round_mode=mode).apply(1195) == expected


def test_write_quote_formulas(tmp_path):
    items = price_items([LineItem("A", 2, "個", 1000), LineItem("B", 1, "式", 5000)], PricingRule(rate=0.9))
    out = write_quote(tmp_path / "q.xlsx", items, QuoteInfo(customer="テスト株式会社", issue_date=date(2026, 9, 29)))
    wb = load_workbook(out)
    ws = wb["見積書"]
    assert ws["A3"].value.startswith("テスト株式会社")
    assert ws["B16"].value == "A" and ws["F16"].value == 900 and ws["G16"].value == "=D16*F16"
    assert ws["G18"].value == "=SUM(G16:G17)"
    assert ws["G19"].value == "=ROUNDDOWN(G18*0.1,0)"
    assert ws["C12"].value == "=G20"
    assert "他社比較" in wb.sheetnames
    assert wb["他社比較"]["D4"].value == 1000


def test_cli(tmp_path, capsys):
    out = tmp_path / "out.xlsx"
    rc = main([str(SAMPLES / "competitor_quote.xlsx"), "-o", str(out), "--customer", "〇〇工業株式会社",
               "--company", str(tmp_path / "none.json")])
    assert rc == 0 and out.exists()
    assert "明細 5 行" in capsys.readouterr().out
