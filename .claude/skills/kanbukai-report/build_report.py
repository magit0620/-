"""data.json から幹部会用「受注の現状」.docx を作る（2ページ：受注済み／未受注）。
使い方: python3 build_report.py data.json template.docx 出力.docx
data.json の形式は extract.py の出力と同じ。追加キー:
  "date": "2026年10月25日", "prev": "10/6", "new_since": "10/7"
行の "色": 完工 / 受注 / 失注 / 新規 / ""（行全体を塗りつぶす）
行の "赤字": 赤字にする列名のリスト（所在・名前・担当・工事内容・金額・変更）
行の "取消線": true で名前と工事内容に取消線"""
import sys, json, docx
from docx.shared import Pt, RGBColor, Twips, Cm
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

data = json.load(open(sys.argv[1]))
d = docx.Document(sys.argv[2])
body = d.element.body
for el in list(body):
    if el.tag != qn('w:sectPr'):
        body.remove(el)
for sec in d.sections:
    sec.left_margin = sec.right_margin = Cm(1.2); sec.top_margin = sec.bottom_margin = Cm(0.7)

RED = RGBColor(0xC0, 0x00, 0x00)
FILL = {'完工': 'C6EFCE', '受注': 'FFEB9C', '失注': 'D9D9D9', '新規': 'BDD7EE'}
KEYS = ['所在', '名前', '担当', '工事内容', '金額', '変更']
LW = [1000, 2200, 1200, 3200, 3300, 4400]
date, prev, new_since = data['date'], data['prev'], data['new_since']
short = date.split('年', 1)[1].replace('月', '/').replace('日', '')

def shd(pr, fill):
    s = OxmlElement('w:shd'); s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), fill); pr.append(s)

def tight(p):
    pf = p.paragraph_format; pf.space_before = Pt(0); pf.space_after = Pt(0); pf.line_spacing = 1.0
    sg = OxmlElement('w:snapToGrid'); sg.set(qn('w:val'), '0'); p._p.get_or_add_pPr().append(sg)

def para(size=8.5, after=0):
    p = d.add_paragraph(); tight(p); p.paragraph_format.space_after = Pt(after); return p

def run(p, text, size=8.5, bold=False, red=False, fill=None, strike=False):
    r = p.add_run(text); r.font.size = Pt(size); r.bold = bold; r.font.strike = strike
    if red: r.font.color.rgb = RED
    if fill: shd(r._r.get_or_add_rPr(), fill)
    return r

def borders(tbl):
    b = OxmlElement('w:tblBorders')
    for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        x = OxmlElement('w:' + e); x.set(qn('w:val'), 'single'); x.set(qn('w:sz'), '4'); x.set(qn('w:space'), '0'); x.set(qn('w:color'), '000000'); b.append(x)
    tbl._tbl.tblPr.append(b)
    cm = OxmlElement('w:tblCellMar')
    for e, v in (('top', 15), ('bottom', 15), ('left', 50), ('right', 50)):
        x = OxmlElement('w:' + e); x.set(qn('w:w'), str(v)); x.set(qn('w:type'), 'dxa'); cm.append(x)
    tbl._tbl.tblPr.append(cm)

def header(title):
    run(para(), f'受注の現状　{title}（{date}現在）', 12, True)
    p = para(after=4)
    run(p, f'{prev}→{short}の比較（100万円以上）　【色の説明】', 8, True)
    run(p, '赤字', 8, True, red=True); run(p, '＝変更　', 8)
    for k, label in (('完工', '完工・請求'), ('受注', '受注'), ('失注', '失注'), ('新規', f'新規（{new_since}以降）')):
        run(p, f'　{"緑黄灰青"["完工受注失注新規".index(k) // 2]}　', 8, fill=FILL[k]); run(p, f'＝{label}　', 8)
    run(p, '取消線', 8, strike=True); run(p, f'＝{prev}時点で完了・失注　黒字＝変更なし', 8)

def table(rows, sz, exact=None, last=False):
    t = d.add_table(rows=1, cols=6); borders(t); t.autofit = False
    for i, col in enumerate(t.columns): col.width = Twips(LW[i])
    for i, h in enumerate(['物件所在', '名前', '担当', '工事内容', '金額（税抜）', f'{short}時点の変更']):
        c = t.rows[0].cells[i]; c.width = Twips(LW[i]); shd(c._tc.get_or_add_tcPr(), 'F2F2F2')
        tight(c.paragraphs[0]); run(c.paragraphs[0], h, sz, True)
    for rec in rows:
        tr = t.add_row()
        if exact:
            h = OxmlElement('w:trHeight'); h.set(qn('w:val'), str(exact)); h.set(qn('w:hRule'), 'exact'); tr._tr.get_or_add_trPr().append(h)
        color = rec.get('色', '')
        for i, k in enumerate(KEYS):
            c = tr.cells[i]; c.width = Twips(LW[i]); tight(c.paragraphs[0])
            if color: shd(c._tc.get_or_add_tcPr(), FILL[color])
            run(c.paragraphs[0], rec.get(k, ''), sz, bold=(k == '変更' and bool(color)),
                red=k in rec.get('赤字', []), strike=rec.get('取消線', False) and k in ('名前', '工事内容'))
    if not last: para(4)

def section_title(text, sz):
    run(para(), text, sz, True)

secs = data['sections']
page1 = [s for s in secs if s['title'].startswith('【受注済み')]
page2 = [s for s in secs if not s['title'].startswith('【受注済み')]
header('受注済み案件')
for s in page1:
    section_title(s['title'], 9.5); table(s['rows'], 8.5)
if data.get('notes'):
    run(para(), '◆その他（要確認）', 10, True)
    for n in data['notes']:
        run(para(), '・' + n, 8.5)
para().add_run().add_break(WD_BREAK.PAGE)
header('未受注（見込み）案件')
n2 = sum(len(s['rows']) for s in page2)
sz2, ex2 = (7, 190) if n2 > 30 else (8.5, None)
for i, s in enumerate(page2):
    section_title(s['title'], 9.5); table(s['rows'], sz2, ex2, last=(i == len(page2) - 1))
d.save(sys.argv[3])
