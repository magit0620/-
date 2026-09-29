"""他社見積ファイル (xlsx / csv / pdf) から明細行を読み取る."""

from __future__ import annotations

import csv
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Optional

# 列見出しの候補 (正規化後の文字列に含まれていればその列とみなす)
HEADER_KEYWORDS = {
    "name": ("品名", "品目", "項目", "内容", "摘要", "名称", "商品名", "作業内容", "明細"),
    "spec": ("仕様", "規格", "型番", "型式", "品番"),
    "quantity": ("数量", "数", "qty"),
    "unit": ("単位",),
    "unit_price": ("単価",),
    "amount": ("金額", "価格", "小計金額", "合計金額"),
}

# 明細ではなく集計行とみなす品名
SUMMARY_WORDS = ("小計", "合計", "消費税", "税込", "税抜", "総額", "値引後", "計")


@dataclass
class LineItem:
    name: str
    quantity: float = 1.0
    unit: str = "式"
    unit_price: float = 0.0
    spec: str = ""

    @property
    def amount(self) -> float:
        return self.quantity * self.unit_price


@dataclass
class CompetitorQuote:
    items: list[LineItem]
    source: str = ""
    vendor: str = ""


def normalize(text: object) -> str:
    """全角英数・記号を半角化し、空白を除去する."""
    if text is None:
        return ""
    s = unicodedata.normalize("NFKC", str(text))
    return re.sub(r"\s+", "", s)


def parse_number(value: object) -> Optional[float]:
    """'¥1,200' / '１２００円' / '△500' などを数値にする. 数値でなければ None."""
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = normalize(value)
    if not s:
        return None
    negative = s.startswith(("△", "▲", "-")) or (s.startswith("(") and s.endswith(")"))
    s = re.sub(r"[¥￥$円,、()△▲\-]", "", s)
    if not re.fullmatch(r"\d+(\.\d+)?", s):
        return None
    n = float(s)
    return -n if negative else n


def _match_header(cell: object) -> Optional[str]:
    text = normalize(cell).lower()
    if not text:
        return None
    # 「単価」を「数」より先に判定するため、より具体的な列から照合する
    for key in ("unit_price", "amount", "unit", "spec", "quantity", "name"):
        for kw in HEADER_KEYWORDS[key]:
            if key == "quantity" and kw == "数":
                if text in ("数", "数量"):
                    return key
                continue
            if kw in text:
                return key
    return None


def detect_columns(rows: list[list[object]]) -> tuple[int, dict[str, int]]:
    """見出し行を探し、(見出し行インデックス, {項目: 列番号}) を返す."""
    best: tuple[int, dict[str, int]] = (-1, {})
    for i, row in enumerate(rows[:50]):
        cols: dict[str, int] = {}
        for j, cell in enumerate(row):
            key = _match_header(cell)
            if key and key not in cols:
                cols[key] = j
        score = len(cols)
        if "name" in cols and ("unit_price" in cols or "amount" in cols) and score > len(best[1]):
            best = (i, cols)
    if best[0] < 0:
        raise ValueError("見積明細の見出し行 (品名・単価/金額など) が見つかりませんでした")
    return best


def _cell(row: list[object], idx: Optional[int]) -> object:
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def extract_items(rows: list[list[object]]) -> list[LineItem]:
    header_idx, cols = detect_columns(rows)
    items: list[LineItem] = []
    blank_streak = 0
    for row in rows[header_idx + 1:]:
        name = str(_cell(row, cols.get("name")) or "").strip()
        qty = parse_number(_cell(row, cols.get("quantity")))
        price = parse_number(_cell(row, cols.get("unit_price")))
        amount = parse_number(_cell(row, cols.get("amount")))

        if not name and price is None and amount is None:
            blank_streak += 1
            if blank_streak >= 5 and items:
                break
            continue
        blank_streak = 0

        norm_name = normalize(name)
        if not norm_name or any(norm_name.startswith(w) for w in SUMMARY_WORDS):
            continue
        if price is None and amount is None:
            continue  # 見出しの続きや備考行

        if qty is None or qty == 0:
            qty = 1.0
        if price is None:
            price = amount / qty  # 金額のみ記載の場合は単価を逆算
        unit = str(_cell(row, cols.get("unit")) or "").strip() or "式"
        spec = str(_cell(row, cols.get("spec")) or "").strip()
        items.append(LineItem(name=name, quantity=qty, unit=unit, unit_price=price, spec=spec))
    if not items:
        raise ValueError("見積明細が1行も読み取れませんでした")
    return items


def _find_vendor(rows: Iterable[list[object]]) -> str:
    pattern = re.compile(r"(株式会社|有限会社|合同会社|\(株\)|\(有\))")
    for row in list(rows)[:15]:
        for cell in row:
            text = unicodedata.normalize("NFKC", str(cell or "")).strip()
            if pattern.search(text) and "御中" not in text and "様" not in text:
                return text
    return ""


def read_rows(path: Path) -> list[list[object]]:
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(path, data_only=True, read_only=True)
        # 明細が読み取れるシートを優先
        sheets = [[list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets]
        for rows in sheets:
            try:
                detect_columns(rows)
                return rows
            except ValueError:
                continue
        return sheets[0] if sheets else []
    if suffix in (".csv", ".tsv", ".txt"):
        raw = path.read_bytes()
        for enc in ("utf-8-sig", "cp932"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ValueError(f"文字コードを判別できません: {path}")
        delim = "\t" if suffix == ".tsv" else ","
        return [list(r) for r in csv.reader(text.splitlines(), delimiter=delim)]
    if suffix == ".pdf":
        import pdfplumber

        rows: list[list[object]] = []
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                for table in page.extract_tables():
                    rows.extend(table)
                if not rows:
                    # 罫線のない PDF はテキストを空白区切りで分割
                    for line in (page.extract_text() or "").splitlines():
                        rows.append(re.split(r"\s{1,}", line.strip()))
        return rows
    raise ValueError(f"未対応のファイル形式です: {suffix} (xlsx / csv / pdf に対応)")


def read_competitor_quote(path: str | Path) -> CompetitorQuote:
    p = Path(path)
    rows = read_rows(p)
    return CompetitorQuote(items=extract_items(rows), source=p.name, vendor=_find_vendor(rows))
