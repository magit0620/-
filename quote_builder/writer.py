"""見積書 Excel を作成する."""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins

from .pricing import PricedItem

FONT = "游ゴシック"
YEN = '"¥"#,##0'
NUM = "#,##0.##"
THIN = Side(style="thin", color="808080")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="DDE5F0")


@dataclass
class QuoteInfo:
    customer: str = ""
    subject: str = ""
    quote_no: str = ""
    issue_date: date = field(default_factory=date.today)
    valid_days: int = 30
    tax_rate: float = 0.10
    delivery: str = "別途ご相談"
    payment: str = "月末締め翌月末払い"
    notes: str = ""
    company: dict = field(default_factory=dict)


def _f(size: int = 10, bold: bool = False) -> Font:
    return Font(name=FONT, size=size, bold=bold)


def _write_quote_sheet(ws, items: list[PricedItem], info: QuoteInfo) -> None:
    ws.title = "見積書"
    widths = {"A": 5, "B": 34, "C": 16, "D": 8, "E": 7, "F": 13, "G": 15}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    ws.merge_cells("A1:G1")
    ws["A1"] = "御 見 積 書"
    ws["A1"].font = _f(20, True)
    ws["A1"].alignment = Alignment(horizontal="center")

    ws["F3"], ws["G3"] = "見積番号", info.quote_no
    ws["F4"], ws["G4"] = "発行日", info.issue_date
    ws["G4"].number_format = "yyyy年m月d日"
    for r in (3, 4):
        ws[f"F{r}"].font = _f()
        ws[f"G{r}"].font = _f()
        ws[f"G{r}"].alignment = Alignment(horizontal="right")

    ws.merge_cells("A3:D3")
    ws["A3"] = f"{info.customer}　御中" if info.customer else "　　　　　　　　　御中"
    ws["A3"].font = _f(14, True)
    ws["A3"].border = Border(bottom=Side(style="medium"))

    ws["A5"] = "下記のとおりお見積り申し上げます。"
    ws["A5"].font = _f()

    c = info.company
    company_lines = [
        c.get("name", ""),
        c.get("postal", "") and f"〒{c['postal']}",
        c.get("address", ""),
        c.get("tel", "") and f"TEL: {c['tel']}",
        c.get("email", ""),
        c.get("person", "") and f"担当: {c['person']}",
        c.get("registration_no", "") and f"登録番号: {c['registration_no']}",
    ]
    r = 6
    for line in [x for x in company_lines if x]:
        ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=7)
        cell = ws.cell(row=r, column=5, value=line)
        cell.font = _f(11 if r == 6 else 9, r == 6)
        r += 1

    labels = [
        ("件名", info.subject),
        ("納期", info.delivery),
        ("支払条件", info.payment),
        ("有効期限", info.issue_date + timedelta(days=info.valid_days)),
    ]
    for i, (label, value) in enumerate(labels):
        row = 7 + i
        ws.cell(row=row, column=1, value=label).font = _f()
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
        v = ws.cell(row=row, column=2, value=value)
        v.font = _f()
        v.border = Border(bottom=THIN)
        if isinstance(value, date):
            v.number_format = "yyyy年m月d日"
            v.alignment = Alignment(horizontal="left")

    header_row = 15
    first = header_row + 1
    last = first + max(len(items), 1) - 1
    subtotal_row = last + 1
    tax_row = subtotal_row + 1
    total_row = tax_row + 1

    ws["A12"] = "御見積金額(税込)"
    ws["A12"].font = _f(12, True)
    ws.merge_cells("A12:B13")
    ws.merge_cells("C12:D13")
    ws["C12"] = f"=G{total_row}"
    ws["C12"].number_format = YEN + "-"
    ws["C12"].font = _f(18, True)
    ws["C12"].alignment = Alignment(horizontal="right", vertical="center")
    ws["A12"].alignment = Alignment(vertical="center")
    for col in "ABCD":
        for rr in (12, 13):
            ws[f"{col}{rr}"].border = Border(bottom=Side(style="double")) if rr == 13 else Border()

    headers = ["No.", "品名", "仕様・型番", "数量", "単位", "単価", "金額"]
    for j, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=j, value=h)
        cell.font = _f(10, True)
        cell.fill = HEAD_FILL
        cell.border = BOX
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for i, p in enumerate(items):
        row = first + i
        it = p.item
        values = [i + 1, it.name, it.spec, it.quantity, it.unit, it.unit_price, f"=D{row}*F{row}"]
        for j, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=j, value=v)
            cell.font = _f()
            cell.border = BOX
        ws.cell(row=row, column=1).alignment = Alignment(horizontal="center")
        ws.cell(row=row, column=2).alignment = Alignment(wrap_text=True, vertical="center")
        ws.cell(row=row, column=4).number_format = NUM
        ws.cell(row=row, column=5).alignment = Alignment(horizontal="center")
        ws.cell(row=row, column=6).number_format = YEN
        ws.cell(row=row, column=7).number_format = YEN

    tax_pct = int(round(info.tax_rate * 100))
    totals = [
        (subtotal_row, "小計", f"=SUM(G{first}:G{last})"),
        (tax_row, f"消費税({tax_pct}%)", f"=ROUNDDOWN(G{subtotal_row}*{info.tax_rate},0)"),
        (total_row, "合計", f"=G{subtotal_row}+G{tax_row}"),
    ]
    for row, label, formula in totals:
        ws.merge_cells(start_row=row, start_column=5, end_row=row, end_column=6)
        lc = ws.cell(row=row, column=5, value=label)
        lc.font = _f(10, True)
        lc.alignment = Alignment(horizontal="center")
        lc.fill = HEAD_FILL
        for col in (5, 6, 7):
            ws.cell(row=row, column=col).border = BOX
        vc = ws.cell(row=row, column=7, value=formula)
        vc.number_format = YEN
        vc.font = _f(10, row == total_row)

    notes_row = total_row + 2
    ws.cell(row=notes_row, column=1, value="備考").font = _f(10, True)
    ws.merge_cells(start_row=notes_row + 1, start_column=1, end_row=notes_row + 4, end_column=7)
    note = ws.cell(row=notes_row + 1, column=1, value=info.notes)
    note.font = _f()
    note.alignment = Alignment(wrap_text=True, vertical="top")
    for rr in range(notes_row + 1, notes_row + 5):
        for cc in range(1, 8):
            ws.cell(row=rr, column=cc).border = BOX

    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.6, bottom=0.6)
    ws.print_area = f"A1:G{notes_row + 4}"
    ws.print_title_rows = f"{header_row}:{header_row}"


def _write_compare_sheet(wb: Workbook, items: list[PricedItem], vendor: str, source: str) -> None:
    ws = wb.create_sheet("他社比較")
    ws["A1"] = f"他社見積比較  (比較元: {vendor or '他社'} / {source})"
    ws["A1"].font = _f(12, True)
    headers = ["No.", "品名", "数量", "他社単価", "自社単価", "単価差", "他社金額", "自社金額", "差額", "率"]
    widths = [5, 34, 8, 12, 12, 12, 14, 14, 14, 8]
    for j, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=3, column=j, value=h)
        cell.font = _f(10, True)
        cell.fill = HEAD_FILL
        cell.border = BOX
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[cell.column_letter].width = w

    first = 4
    for i, p in enumerate(items):
        r = first + i
        row = [
            i + 1, p.item.name, p.item.quantity, p.competitor_price, p.item.unit_price,
            f"=E{r}-D{r}", f"=C{r}*D{r}", f"=C{r}*E{r}", f"=H{r}-G{r}", f"=IF(D{r}=0,\"\",E{r}/D{r})",
        ]
        for j, v in enumerate(row, start=1):
            cell = ws.cell(row=r, column=j, value=v)
            cell.font = _f()
            cell.border = BOX
            if j in (4, 5, 6, 7, 8, 9):
                cell.number_format = YEN + ';[Red]-"¥"#,##0'
            elif j == 10:
                cell.number_format = "0.0%"
            elif j == 3:
                cell.number_format = NUM
    last = first + max(len(items), 1) - 1
    tr = last + 1
    ws.cell(row=tr, column=2, value="合計(税抜)").font = _f(10, True)
    for col in "GHI":
        c = ws[f"{col}{tr}"]
        c.value = f"=SUM({col}{first}:{col}{last})"
        c.number_format = YEN + ';[Red]-"¥"#,##0'
        c.font = _f(10, True)
    ws[f"J{tr}"] = f'=IF(G{tr}=0,"",H{tr}/G{tr})'
    ws[f"J{tr}"].number_format = "0.0%"
    for cc in range(1, 11):
        ws.cell(row=tr, column=cc).border = BOX
    ws.freeze_panes = "C4"


def write_quote(path: str | Path, items: list[PricedItem], info: QuoteInfo,
                vendor: str = "", source: str = "", compare: bool = True) -> Path:
    wb = Workbook()
    _write_quote_sheet(wb.active, items, info)
    if compare:
        _write_compare_sheet(wb, items, vendor, source)
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
